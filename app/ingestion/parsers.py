import json
from abc import ABC, abstractmethod
from pathlib import Path

from docx import Document as DocxFile
from pypdf import PdfReader

from app.ingestion.models import Document


class ParseError(Exception):
    """ფაილის წაკითხვა ვერ მოხერხდა."""


class BaseParser(ABC):
    extensions: tuple[str, ...] = ()

    @abstractmethod
    def parse(self, path: Path) -> list[Document]:
        """ერთი ფაილიდან ვაბრუნებთ ერთ ან რამდენიმე დოკუმენტს."""


class JsonlParser(BaseParser):
    extensions = (".jsonl",)

    def parse(self, path: Path) -> list[Document]:
        docs = []
        with path.open(encoding="utf-8") as f:
            for line_no, line in enumerate(f, start=1):
                line = line.strip()
                if not line:
                    continue
                try:
                    row = json.loads(line)
                    docs.append(
                        Document(
                            doc_id=row["title"],
                            title=row["title"],
                            text=row["text"],
                            source=path.name,
                            format="jsonl",
                            metadata={"line": line_no},
                        )
                    )
                except (json.JSONDecodeError, KeyError) as e:
                    raise ParseError(f"{path.name}, ხაზი {line_no}: {e}") from e
        return docs


class TextParser(BaseParser):
    """Markdown და TXT."""
    extensions = (".md", ".txt")

    def parse(self, path: Path) -> list[Document]:
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError as e:
            raise ParseError(f"{path.name}: არასწორი კოდირება") from e

        title = path.stem
        for line in text.splitlines():
            if line.startswith("# "):          # პირველი # სათაური
                title = line[2:].strip()
                break

        return [Document(
            doc_id=path.stem, title=title, text=text,
            source=path.name, format=path.suffix.lstrip("."),
        )]


class PdfParser(BaseParser):
    extensions = (".pdf",)

    def parse(self, path: Path) -> list[Document]:
        try:
            reader = PdfReader(str(path))
            pages = [(p.extract_text() or "") for p in reader.pages]
        except Exception as e:
            raise ParseError(f"{path.name}: PDF ვერ წაიკითხა ({e})") from e

        meta = reader.metadata
        title = (meta.title if meta and meta.title else path.stem)
        return [Document(
            doc_id=path.stem, title=title, text="\n\n".join(pages),
            source=path.name, format="pdf",
            metadata={"page_count": len(pages)},
        )]


class DocxParser(BaseParser):
    extensions = (".docx",)

    def parse(self, path: Path) -> list[Document]:
        try:
            doc = DocxFile(str(path))
        except Exception as e:
            raise ParseError(f"{path.name}: DOCX ვერ წაიკითხა ({e})") from e

        lines = []
        for p in doc.paragraphs:
            if not p.text.strip():
                continue
            if p.style.name.startswith("Heading"):
                lines.append(f"# {p.text.strip()}")   # სათაურებს ვინარჩუნებთ
            else:
                lines.append(p.text.strip())

                first_heading = next(
            (p.text.strip() for p in doc.paragraphs
             if p.style.name.startswith("Heading") and p.text.strip()),
            None,
        )
        title = doc.core_properties.title or first_heading or path.stem
        return [Document(
            doc_id=path.stem, title=title, text="\n\n".join(lines),
            source=path.name, format="docx",
            metadata={"author": doc.core_properties.author or ""},
        )]


_PARSERS: list[BaseParser] = [JsonlParser(), TextParser(), PdfParser(), DocxParser()]


def parse_file(path: str | Path) -> list[Document]:
    """შესასვლელი წერტილი: ფორმატს გაფართოებით ვარჩევთ."""
    path = Path(path)
    for parser in _PARSERS:
        if path.suffix.lower() in parser.extensions:
            return parser.parse(path)
    raise ParseError(f"მხარდაუჭერელი ფორმატი: {path.suffix}")